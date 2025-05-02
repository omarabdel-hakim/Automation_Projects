import openpyxl
book = openpyxl.load_workbook("C:\\Users\\Mohamed\\Documents\\pythondemo.xlsx")
sheet = book.active
Dich = {}
for Row in range(1,sheet.max_row + 1):

    if sheet.cell(row=Row , column=1).value == "omar abdelhakim":

        for Col in range(1,sheet.max_column + 1):

            Dich[sheet.cell(row=1 , column=Col).value] = sheet.cell(row=Row , column=Col).value
print(Dich)