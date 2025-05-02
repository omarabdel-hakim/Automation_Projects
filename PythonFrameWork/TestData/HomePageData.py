import openpyxl

class HomePageData:
    test_HomePage_data = [{"Name":"omar abdelhakim","Email":"omar@gmail.com","Password":"01153330950"
                                ,"Gender":"Male","Date":"2001/01/01"},
                            {"Name":"sara walid","Email":"sara@gmail.com","Password":"01001212774"
                                ,"Gender":"Female","Date":"2002/5/14"}]

    @staticmethod
    def getTestData(test_case_name):
        book = openpyxl.load_workbook("C:\\Users\\Mohamed\\Documents\\pythondemo.xlsx")
        sheet = book.active
        Dich = {}
        for Row in range(1, sheet.max_row + 1):

            if sheet.cell(row=Row, column=1).value == test_case_name:

                for Col in range(1, sheet.max_column + 1):
                    Dich[sheet.cell(row=1, column=Col).value] = sheet.cell(row=Row, column=Col).value
        return [Dich]